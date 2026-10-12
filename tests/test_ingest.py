from PIL import Image

from engine.ingest import Ingestor


def _make_docx(path):
    import docx

    d = docx.Document()
    d.add_heading("Quarterly results", 1)
    d.add_paragraph("Revenue grew 18 percent.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Region", "Sales"
    t.cell(1, 0).text, t.cell(1, 1).text = "North", "1.2M"
    img = path.parent / "chart.png"
    Image.new("RGB", (400, 300), (200, 30, 30)).save(img)
    d.add_picture(str(img))
    d.save(path)
    img.unlink()


def _make_pptx(path):
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Old deck title"
    s.placeholders[1].text = "An old bullet"
    prs.save(path)


def test_ingest_many_formats(tmp_path, cfg):
    src = tmp_path / "src"
    src.mkdir()
    (src / "notes.md").write_text("# Notes\n\nCustomer churn fell to 2.9%.", encoding="utf-8")
    (src / "data.csv").write_text("month,revenue\nJul,1.2\nAug,1.4\n", encoding="utf-8")
    _make_docx(src / "report.docx")
    _make_pptx(src / "old.pptx")
    Image.new("RGB", (640, 480), (20, 120, 60)).save(src / "photo.jpg")
    Image.new("RGB", (16, 16)).save(src / "icon.png")  # too small -> kept as-is, not extracted

    ing = Ingestor(cfg, llm=None).run(sorted(src.iterdir()), [], "Board update", tmp_path / "out")
    md = ing.markdown
    for needle in ("Board update", "churn fell", "Quarterly results", "| Region | Sales |",
                   "Old deck title", "An old bullet", "Jul"):
        assert needle in md, needle
    names = [a.name for a in ing.assets]
    assert any(n.startswith("report-img") for n in names)  # image pulled out of the docx
    assert any(n.startswith("photo") for n in names)
    assert (tmp_path / "out" / "source.md").exists()
    assert not ing.warnings
