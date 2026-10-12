import re

from PIL import Image

from engine.brand import finalize_brand
from engine.render import render_deck
from engine.schema import Deck


def test_html_deck_renders_every_slide(tmp_path, sample_deck):
    brand = finalize_brand({"name": "t"})
    outputs, warnings = render_deck(sample_deck, brand, tmp_path / "tpl", tmp_path / "v001", formats=("html",))
    html = outputs["html"].read_text(encoding="utf-8")
    assert html.count('<section class="slide') == len(sample_deck.slides)
    assert "window.deck" in html and "@keyframes" in html
    assert 'class="chart"' in html and "data-count=" in html
    assert re.search(r'class="pie"', html)  # doughnut chart
    assert not [w for w in warnings if "failed" in w]


def test_pptx_export(tmp_path, sample_deck):
    from pptx import Presentation

    brand = finalize_brand({"name": "t"})
    outputs, warnings = render_deck(sample_deck, brand, tmp_path / "tpl", tmp_path / "v001", formats=("html", "pptx"))
    prs = Presentation(str(outputs["pptx"]))
    assert len(prs.slides) == len(sample_deck.slides)
    assert not [w for w in warnings if "failed" in w]


def test_images_and_missing_images(tmp_path):
    assets = tmp_path / "assets"
    assets.mkdir()
    Image.new("RGB", (800, 500), (30, 90, 200)).save(assets / "team.jpg")
    deck = Deck.model_validate({"title": "x", "slides": [
        {"layout": "image", "title": "Our team", "image": "team.jpg"},
        {"layout": "image_text", "title": "Office", "image": "team.png", "bullets": ["a"]},  # wrong ext tolerated
        {"layout": "image", "title": "Missing", "image": "nope.png"},
    ]})
    outputs, warnings = render_deck(deck, finalize_brand({}), tmp_path / "tpl", tmp_path / "v1",
                                    asset_dirs=[assets], formats=("html", "pptx"))
    html = outputs["html"].read_text(encoding="utf-8")
    assert html.count("data:image/jpeg;base64") >= 2
    assert any("nope.png" in w for w in warnings)


def test_brand_colors_and_dark_background(tmp_path, sample_deck):
    brand = finalize_brand({"colors": {"background": "#0B1020", "text": "#0B1020", "primary": "#22D3EE"}})
    assert brand["colors"]["text"] != "#0B1020"  # fixed unreadable text colour
    outputs, _ = render_deck(sample_deck, brand, tmp_path / "tpl", tmp_path / "v1")
    assert "--bg:#0B1020" in outputs["html"].read_text(encoding="utf-8")
