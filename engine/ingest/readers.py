"""One reader per input format. Each returns Markdown and saves embedded images to assets_dir."""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Callable

from ..util import log, md_table, slugify

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".aac", ".mp4", ".mov", ".mkv", ".webm"}
TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".json", ".yaml", ".yml", ".xml", ".log"}
SHEET_EXT = {".xlsx", ".xlsm", ".xls", ".csv", ".tsv"}
HTML_EXT = {".html", ".htm"}
SUPPORTED = IMAGE_EXT | AUDIO_EXT | TEXT_EXT | SHEET_EXT | HTML_EXT | {".pdf", ".docx", ".pptx", ".url"}


class ImageSaver:
    """Normalises any image blob to PNG/JPEG under assets_dir, skipping tiny ones."""

    def __init__(self, assets_dir: Path, min_px: int = 120, limit: int = 30):
        self.dir = assets_dir
        self.min_px = min_px
        self.limit = limit
        self.saved: list[Path] = []

    def save(self, data: bytes | Path, stem: str) -> Path | None:
        if len(self.saved) >= self.limit:
            return None
        from PIL import Image

        try:
            src = Image.open(io.BytesIO(data) if isinstance(data, (bytes, bytearray)) else data)
            src.load()
        except Exception:
            return None  # EMF/WMF/corrupt - skip
        if min(src.size) < self.min_px:
            return None
        self.dir.mkdir(parents=True, exist_ok=True)
        alpha = src.mode in ("RGBA", "LA", "P")
        ext = ".png" if alpha else ".jpg"
        base = slugify(stem, 40)
        path = self.dir / f"{base}{ext}"
        k = 2
        while path.exists():
            path = self.dir / f"{base}-{k}{ext}"
            k += 1
        im = src.convert("RGBA" if alpha else "RGB")
        im.thumbnail((2400, 2400))
        im.save(path, "PNG" if alpha else "JPEG", **({} if alpha else {"quality": 90}))
        self.saved.append(path)
        return path


def read_text_file(path: Path) -> str:
    for enc in ("utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return ""


def read_docx(path: Path, saver: ImageSaver | None) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    out: list[str] = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            low = style.lower()
            if low == "title":
                out.append("## " + text)
            elif low.startswith("heading"):
                m = re.search(r"\d+", style)
                level = int(m.group(0)) if m else 1
                out.append("#" * min(level + 1, 6) + " " + text)
            elif "list" in low:
                out.append("- " + text)
            else:
                out.append(text)
        elif tag == "tbl":
            t = Table(child, d)
            out.append(md_table([[c.text for c in r.cells] for r in t.rows]))
    if saver:
        for rel in d.part.rels.values():
            if "image" in rel.reltype:
                try:
                    saver.save(rel.target_part.blob, f"{path.stem}-img")
                except Exception:
                    pass
    return "\n\n".join(out)


def read_pdf(path: Path, saver: ImageSaver | None, ocr: Callable[[Path], str] | None = None,
             ocr_max_pages: int = 30, tmp_dir: Path | None = None) -> str:
    import fitz

    with fitz.open(str(path)) as doc:
        try:
            import pymupdf4llm

            md = pymupdf4llm.to_markdown(doc)
        except Exception:
            md = "\n\n".join(page.get_text() for page in doc)
        texts = [page.get_text().strip() for page in doc]
        empty = [i for i, t in enumerate(texts) if len(t) < 30]
        if ocr and empty and len(empty) >= max(1, len(texts) // 2):
            log.info("%s looks scanned: OCR on %d page(s)", path.name, min(len(empty), ocr_max_pages))
            tmp_dir = tmp_dir or path.parent
            tmp_dir.mkdir(parents=True, exist_ok=True)
            pages = []
            for i in empty[:ocr_max_pages]:
                p = tmp_dir / f"_ocr-{slugify(path.stem)}-{i + 1}.png"
                doc[i].get_pixmap(dpi=150).save(str(p))
                try:
                    pages.append(f"### Page {i + 1}\n\n{ocr(p)}")
                except Exception as e:
                    log.warning("OCR failed on page %d: %s", i + 1, e)
                finally:
                    p.unlink(missing_ok=True)
            md = (md + "\n\n" + "\n\n".join(pages)).strip()
        if saver:
            seen = set()
            for page in doc:
                for img in page.get_images(full=True):
                    xref = img[0]
                    if xref in seen:
                        continue
                    seen.add(xref)
                    try:
                        info = doc.extract_image(xref)
                        saver.save(info["image"], f"{path.stem}-p{page.number + 1}")
                    except Exception:
                        pass
    return md


def read_pptx(path: Path, saver: ImageSaver | None) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        out.append(f"## Slide {i}: {title}".rstrip(": "))
        for shp in slide.shapes:
            if shp == slide.shapes.title:
                continue
            if shp.has_text_frame:
                for p in shp.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        out.append(("  " * p.level) + "- " + t)
            if getattr(shp, "has_table", False) and shp.has_table:
                out.append(md_table([[c.text for c in r.cells] for r in shp.table.rows]))
            if getattr(shp, "has_chart", False) and shp.has_chart:
                ch = shp.chart
                try:
                    cats = list(ch.plots[0].categories)
                    rows = [["series"] + [str(c) for c in cats]]
                    for s in ch.plots[0].series:
                        rows.append([s.name] + [str(v) for v in s.values])
                    out.append("Chart data:\n" + md_table(rows))
                except Exception:
                    pass
            if saver and shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    saver.save(shp.image.blob, f"{path.stem}-s{i}")
                except Exception:
                    pass
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"Notes: {notes}")
    return "\n".join(out)


def read_sheet(path: Path, max_rows: int = 60) -> str:
    import pandas as pd

    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv"):
        frames = {path.stem: pd.read_csv(path, sep="\t" if suffix == ".tsv" else None, engine="python")}
    else:
        frames = pd.read_excel(path, sheet_name=None)
    out = []
    for name, df in frames.items():
        df = df.dropna(how="all").dropna(axis=1, how="all")
        if df.empty:
            continue
        note = f" (first {max_rows} of {len(df)} rows)" if len(df) > max_rows else ""
        out.append(f"### Sheet: {name}{note}\n\n" + df.head(max_rows).to_markdown(index=False))
    return "\n\n".join(out)


def read_html(html: str) -> str:
    import trafilatura

    text = None
    try:
        text = trafilatura.extract(html, output_format="markdown", include_tables=True)
    except TypeError:
        text = None
    return text or trafilatura.extract(html, include_tables=True) or ""


def read_url(url: str) -> str:
    import trafilatura

    downloaded = trafilatura.fetch_url(url)
    if not downloaded:
        raise RuntimeError(f"could not download {url}")
    return read_html(downloaded)


_whisper = None


def transcribe(path: Path, model_size: str = "large-v3") -> str:
    """GPU speech-to-text (optional dependency faster-whisper)."""
    global _whisper
    try:
        from faster_whisper import WhisperModel
    except ImportError as e:
        raise RuntimeError("audio/video input needs `pip install faster-whisper`") from e
    if _whisper is None:
        try:
            _whisper = WhisperModel(model_size, device="cuda", compute_type="float16")
        except Exception as e:
            log.warning("whisper on CUDA failed (%s), using CPU", e)
            _whisper = WhisperModel(model_size, device="cpu", compute_type="int8")
    segments, _ = _whisper.transcribe(str(path), vad_filter=True)
    return " ".join(s.text.strip() for s in segments)


def read_text_only(path: Path) -> str:
    """Plain text from a document without images/OCR (used for brand style guides)."""
    s = path.suffix.lower()
    try:
        if s == ".pdf":
            return read_pdf(path, None)
        if s == ".docx":
            return read_docx(path, None)
        if s in TEXT_EXT:
            return read_text_file(path)
    except Exception as e:
        log.warning("could not read %s: %s", path.name, e)
    return ""
