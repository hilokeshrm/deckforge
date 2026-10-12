"""Optional export: deck JSON -> editable .pptx with the same brand (no animations)."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from ..brand.colors import is_dark, mix, text_on
from ..schema import Deck, Slide
from .common import AssetResolver, fit_size

CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
M = 0.6  # side margin, inches


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


class PptxRenderer:
    def __init__(self, brand: dict, brand_dir: Path, asset_dirs: list[Path]):
        self.b = brand
        self.c = brand["colors"]
        self.dir = Path(brand_dir)
        self.assets = AssetResolver(list(asset_dirs) + [self.dir / "assets" / "images", self.dir / "assets"])
        self.warnings: list[str] = []
        self.base = self.dir / brand["base_pptx"] if brand.get("base_pptx") else None
        if self.base and not self.base.exists():
            self.base = None

    def warn(self, idx, msg):
        self.warnings.append(f"slide {idx} (pptx): {msg}")

    # ---------- primitives ----------
    def _new_prs(self):
        if self.base:
            prs = Presentation(str(self.base))
            lst = prs.slides._sldIdLst
            for sld in list(lst):
                prs.part.drop_rel(sld.rId)
                lst.remove(sld)
        else:
            prs = Presentation()
            prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        return prs

    def _bg(self, slide, color: str, force: bool = False):
        if self.base and self.b.get("use_master_background", True) and not force:
            return
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(color)

    def _rect(self, slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, alpha=None):
        shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(color)
        shp.line.fill.background()
        shp.shadow.inherit = False
        if alpha is not None:
            srgb = shp._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            a = OxmlElement("a:alpha")
            a.set("val", str(int(alpha * 100000)))
            srgb.append(a)
        return shp

    def _font(self, font, size, color, bold=False, italic=False, heading=False):
        font.size = Pt(size)
        font.bold = bold
        font.italic = italic
        font.name = self.b["fonts"]["heading" if heading else "body"]
        font.color.rgb = rgb(color)

    def _text(self, slide, idx, x, y, w, h, text, size, color, bold=False, italic=False, heading=False,
              align="left", anchor="top", min_size=None):
        lines = str(text or "").split("\n")
        fs, ok = fit_size(lines, w * 72, h * 72, size, min_size or max(12, int(size * 0.55)))
        if not ok:
            self.warn(idx, f"text may overflow: {lines[0][:40]}")
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = ANCHOR[anchor]
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Inches(0.04))
        for j, line in enumerate(lines):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.alignment = ALIGN[align]
            p.line_spacing = 1.05
            r = p.add_run()
            r.text = line
            self._font(r.font, fs, color, bold, italic, heading)
        return tb

    def _bullets(self, slide, idx, items, x, y, w, h, size, color, dot):
        items = [str(i).strip() for i in items if str(i).strip()]
        if not items:
            return
        fs, ok = fit_size(items, (w - 0.45) * 72, h * 72, size, 12, gap=0.5)
        if not ok:
            self.warn(idx, "bullets may overflow")
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        for j, item in enumerate(items):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = Pt(fs * 0.5)
            p.line_spacing = 1.05
            clean = item.replace("**", "")
            head, rest = "", clean
            if ": " in clean:
                h_, r_ = clean.split(": ", 1)
                if len(h_) <= 40:
                    head, rest = h_, r_
            if head:
                r = p.add_run()
                r.text = head + ": "
                self._font(r.font, fs, self.c["primary"] if not is_dark(self.c["background"]) else self.c["accent"], bold=True)
            r = p.add_run()
            r.text = rest
            self._font(r.font, fs, color)
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.34))))
            pPr.set("indent", str(-int(Inches(0.34))))
            clr = OxmlElement("a:buClr")
            s = OxmlElement("a:srgbClr")
            s.set("val", dot.lstrip("#"))
            clr.append(s)
            pPr.append(clr)
            bu = OxmlElement("a:buChar")
            bu.set("char", "●")
            pPr.append(bu)

    def _picture(self, slide, path: Path, x, y, w, h, cover=True):
        from PIL import Image

        with Image.open(path) as im:
            iw, ih = im.size
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        if cover:
            box, img = w / h, iw / ih
            if img > box:
                c = (1 - box / img) / 2
                pic.crop_left = pic.crop_right = c
            else:
                c = (1 - img / box) / 2
                pic.crop_top = pic.crop_bottom = c
        return pic

    def _logo(self, slide, dark, big=False, pos=None):
        b = self.b
        rel = (b.get("logo_on_dark") or b.get("logo")) if dark else b.get("logo")
        if not rel or not (self.dir / rel).exists():
            return
        from PIL import Image

        path = self.dir / rel
        with Image.open(path) as im:
            iw, ih = im.size
        h = 0.75 if big else 0.42
        w = min(h * iw / ih, 2.8 if big else 1.9)
        h = w * ih / iw
        pos = pos or b.get("logo_position", "top-right")
        x = self.W - M - w if "right" in pos else M
        y = 0.4 if "top" in pos else self.H - 0.35 - h
        if dark and not b.get("logo_on_dark") and not b.get("logo_has_alpha", True):
            self._rect(slide, x - 0.08, y - 0.06, w + 0.16, h + 0.12, "#FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE)
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))

    def _chrome(self, slide, idx, total, bg):
        dark = is_dark(bg)
        self._logo(slide, dark)
        muted = mix(text_on(bg), bg, 0.45)
        right = self.b.get("logo_position") != "bottom-right"
        self._text(slide, idx, self.W - M - 1.2 if right else M, self.H - 0.5, 1.2, 0.3, f"{idx} / {total}", 11, muted,
                   align="right" if right else "left")
        if self.b.get("footer"):
            self._text(slide, idx, M if right else self.W / 2 - 3, self.H - 0.5, 6, 0.3, self.b["footer"], 11, muted)

    def _title(self, slide, idx, s: Slide, color):
        self._text(slide, idx, M, 0.38, self.W - 2 * M - 2.2, 1.15, s.title, 32, color, bold=True, heading=True,
                   anchor="bottom", min_size=20)
        self._rect(slide, M, 1.6, 0.9, 0.07, self.c["accent"])
        return 1.95

    # ---------- layouts ----------
    def render(self, deck: Deck, out: Path) -> list[str]:
        prs = self._new_prs()
        self.W, self.H = prs.slide_width / 914400, prs.slide_height / 914400
        blank = min(prs.slide_layouts, key=lambda l: len(l.placeholders))
        total, sections = len(deck.slides), 0
        for i, s in enumerate(deck.slides, 1):
            slide = prs.slides.add_slide(blank)
            for ph in list(slide.placeholders):
                ph._element.getparent().remove(ph._element)
            layout = s.layout
            if layout == "section":
                sections += 1
            try:
                getattr(self, f"_l_{layout}", self._l_bullets)(slide, i, total, s, deck, sections)
            except Exception as ex:
                self.warn(i, f"layout {layout} failed ({ex}); rendered as bullets")
                for shp in list(slide.shapes):
                    shp._element.getparent().remove(shp._element)
                self._l_bullets(slide, i, total, s, deck, sections)
            if s.notes:
                slide.notes_slide.notes_text_frame.text = s.notes
        out.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out))
        return self.warnings

    def _l_title(self, slide, i, total, s, deck, _n, closing=False):
        c = self.c
        img = self.dir / self.b["background_image"] if self.b.get("background_image") else None
        if self.b.get("title_background") == "image" and img and img.exists():
            self._bg(slide, "#000000", force=True)
            self._picture(slide, img, 0, 0, self.W, self.H)
            self._rect(slide, 0, 0, self.W, self.H, "#000000", alpha=0.5)
            bg = "#000000"
        else:
            bg = c["background"] if self.b.get("title_background") == "background" and not closing else c["primary"]
            self._bg(slide, bg, force=True)
            self._rect(slide, self.W - 3.2, -1.6, 4.8, 4.8, c["accent"], MSO_SHAPE.OVAL, alpha=0.3)
        fg = text_on(bg)
        self._logo(slide, is_dark(bg), big=True, pos="top-left")
        align = "center" if closing else "left"
        x, w = (M, self.W - 2 * M) if closing else (M, self.W * 0.75)
        title = s.title or (deck.title if not closing else "Thank you")
        self._text(slide, i, x, self.H * 0.3, w, 1.9, title, 44, fg, bold=True, heading=True, anchor="bottom", align=align)
        bx = self.W / 2 - 0.6 if closing else M
        self._rect(slide, bx, self.H * 0.3 + 2.0, 1.2, 0.08, c["accent"])
        sub = s.subtitle or (deck.subtitle if not closing else None)
        if sub:
            self._text(slide, i, x, self.H * 0.3 + 2.2, w, 1.0, sub, 22, mix(fg, bg, 0.25), align=align)
        if closing and s.bullets:
            self._bullets(slide, i, s.bullets, self.W / 2 - 3, self.H * 0.3 + 3.2, 6, 1.6, 16, fg, c["accent"])

    def _l_closing(self, slide, i, total, s, deck, n):
        self._l_title(slide, i, total, s, deck, n, closing=True)

    def _l_section(self, slide, i, total, s, deck, n):
        bg = self.c["surface"]
        self._bg(slide, bg, force=True)
        fg = text_on(bg)
        self._text(slide, i, M, self.H * 0.2, 4, 1.7, f"{n:02d}", 96, self.c["accent"], bold=True, heading=True)
        self._text(slide, i, M, self.H * 0.2 + 1.8, self.W - 2 * M, 1.4, s.title, 40, fg, bold=True, heading=True)
        if s.subtitle:
            self._text(slide, i, M, self.H * 0.2 + 3.2, self.W - 2 * M, 0.8, s.subtitle, 20, mix(fg, bg, 0.4))
        self._chrome(slide, i, total, bg)

    def _l_bullets(self, slide, i, total, s, deck, _n):
        bg = self.c["background"]
        self._bg(slide, bg)
        y = self._title(slide, i, s, self.c["text"])
        if s.subtitle:
            self._text(slide, i, M, y, self.W - 2 * M, 0.5, s.subtitle, 18, self.c["muted"])
            y += 0.6
        img = self.assets.find(s.image)
        items = s.bullets or s.left + s.right
        if img:
            self._bullets(slide, i, items, M, y, self.W * 0.5 - M, self.H - y - 0.8, 22, self.c["text"], self.c["accent"])
            self._picture(slide, img, self.W * 0.55, y, self.W * 0.45 - M, self.H - y - 0.8)
        else:
            self._bullets(slide, i, items, M, y, self.W - 2 * M, self.H - y - 0.8, 22, self.c["text"], self.c["accent"])
        self._chrome(slide, i, total, bg)

    def _l_two_column(self, slide, i, total, s, deck, _n):
        bg = self.c["background"]
        self._bg(slide, bg)
        y = self._title(slide, i, s, self.c["text"]) + 0.1
        left, right = s.left, s.right
        if not left and not right:
            half = (len(s.bullets) + 1) // 2
            left, right = s.bullets[:half], s.bullets[half:]
        cw = (self.W - 2 * M - 0.5) / 2
        for k, (title, items) in enumerate(((s.left_title, left), (s.right_title, right))):
            x = M + k * (cw + 0.5)
            self._rect(slide, x, y, cw, self.H - y - 0.75, self.c["surface"], MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.05
            yy = y + 0.25
            if title:
                self._text(slide, i, x + 0.3, yy, cw - 0.6, 0.6, title, 20, self.c["primary"], bold=True, heading=True)
                yy += 0.7
            self._bullets(slide, i, items, x + 0.3, yy, cw - 0.6, self.H - yy - 1.0, 18, self.c["text"], self.c["accent"])
        self._chrome(slide, i, total, bg)

    def _l_chart(self, slide, i, total, s, deck, n):
        if not s.chart or not s.chart.categories or not s.chart.series:
            self.warn(i, "no chart data")
            return self._l_bullets(slide, i, total, s, deck, n)
        bg = self.c["background"]
        self._bg(slide, bg)
        y = self._title(slide, i, s, self.c["text"])
        spec = s.chart
        cd = CategoryChartData()
        cd.categories = spec.categories
        for ser in spec.series:
            cd.add_series(ser.name, (list(ser.values) + [0] * len(spec.categories))[: len(spec.categories)])
        w = (self.W - 2 * M) * (0.62 if s.bullets else 1)
        gf = slide.shapes.add_chart(CHART_TYPES[spec.type], Inches(M), Inches(y), Inches(w), Inches(self.H - y - 0.7), cd)
        ch = gf.chart
        ch.font.size = Pt(12)
        ch.font.name = self.b["fonts"]["body"]
        ch.font.color.rgb = rgb(self.c["text"])
        palette = self.b["chart_colors"]
        pie = spec.type in ("pie", "doughnut")
        ch.has_legend = pie or len(spec.series) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
        plot = ch.plots[0]
        if pie:
            for k, point in enumerate(plot.series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = rgb(palette[k % len(palette)])
            plot.has_data_labels = True
            plot.data_labels.show_percentage = True
            plot.data_labels.show_value = False
        else:
            plot.vary_by_categories = False
            for k, ser in enumerate(plot.series):
                col = rgb(palette[k % len(palette)])
                if spec.type == "line":
                    ser.format.line.color.rgb = col
                    ser.format.line.width = Pt(3)
                    ser.smooth = False
                else:
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = col
            ch.value_axis.has_major_gridlines = True
            ch.value_axis.major_gridlines.format.line.color.rgb = rgb(mix(bg, self.c["text"], 0.12))
            ch.value_axis.format.line.fill.background()
            if len(spec.series) == 1:
                plot.has_data_labels = True
                plot.data_labels.font.size = Pt(11)
        if s.bullets:
            x = M + w + 0.4
            self._bullets(slide, i, s.bullets, x, y + 0.3, self.W - x - M, self.H - y - 1.2, 18, self.c["text"], self.c["accent"])
        self._chrome(slide, i, total, bg)

    def _l_table(self, slide, i, total, s, deck, n):
        t = s.table
        if not t or not t.header:
            return self._l_bullets(slide, i, total, s, deck, n)
        bg = self.c["background"]
        self._bg(slide, bg)
        y = self._title(slide, i, s, self.c["text"])
        header, rows = t.header[:7], [r[:7] for r in t.rows[:10]]
        nrows, ncols = len(rows) + 1, len(header)
        h = min(0.5 * nrows, self.H - y - 0.8)
        tbl = slide.shapes.add_table(nrows, ncols, Inches(M), Inches(y), Inches(self.W - 2 * M), Inches(h)).table
        size = 14 if nrows <= 7 else 11
        on_p = text_on(self.c["primary"])
        for r in range(nrows):
            vals = header if r == 0 else (rows[r - 1] + [""] * ncols)[:ncols]
            for cidx, v in enumerate(vals):
                cell = tbl.cell(r, cidx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(self.c["primary"] if r == 0 else (self.c["surface"] if r % 2 == 0 else bg))
                cell.text = str(v)
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        self._font(run.font, size, on_p if r == 0 else self.c["text"], bold=(r == 0))
        self._chrome(slide, i, total, bg)

    def _l_stats(self, slide, i, total, s, deck, n):
        stats = s.stats[:4]
        if not stats:
            return self._l_bullets(slide, i, total, s, deck, n)
        bg = self.c["background"]
        self._bg(slide, bg)
        y = self._title(slide, i, s, self.c["text"]) + 0.3
        gap = 0.35
        tw = (self.W - 2 * M - gap * (len(stats) - 1)) / len(stats)
        for k, st in enumerate(stats):
            x = M + k * (tw + gap)
            self._rect(slide, x, y, tw, 3.4, self.c["surface"], MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.08
            self._text(slide, i, x + 0.2, y + 0.5, tw - 0.4, 1.4, st.value, 54, self.c["primary"], bold=True, heading=True,
                       align="center", anchor="middle", min_size=24)
            self._text(slide, i, x + 0.2, y + 2.0, tw - 0.4, 1.1, st.label, 16, self.c["muted"], align="center")
        self._chrome(slide, i, total, bg)

    def _l_image(self, slide, i, total, s, deck, n):
        img = self.assets.find(s.image)
        if not img:
            self.warn(i, f"image '{s.image}' not found")
            return self._l_bullets(slide, i, total, s, deck, n)
        self._picture(slide, img, 0, 0, self.W, self.H)
        self._rect(slide, 0, self.H * 0.6, self.W, self.H * 0.4, "#000000", alpha=0.55)
        self._text(slide, i, M, self.H * 0.64, self.W - 2 * M, 1.1, s.title, 36, "#FFFFFF", bold=True, heading=True)
        cap = s.subtitle or (s.bullets[0] if s.bullets else "")
        if cap:
            self._text(slide, i, M, self.H * 0.64 + 1.15, self.W - 2 * M, 0.7, cap, 18, "#E5E7EB")
        self._chrome(slide, i, total, "#000000")

    def _l_image_text(self, slide, i, total, s, deck, n):
        img = self.assets.find(s.image)
        if not img:
            return self._l_bullets(slide, i, total, s, deck, n)
        bg = self.c["background"]
        self._bg(slide, bg)
        self._picture(slide, img, 0, 0, self.W * 0.45, self.H)
        x = self.W * 0.45 + 0.5
        w = self.W - x - M
        self._text(slide, i, x, 0.6, w, 1.3, s.title, 30, self.c["text"], bold=True, heading=True, anchor="bottom")
        self._rect(slide, x, 2.0, 0.9, 0.07, self.c["accent"])
        self._bullets(slide, i, s.bullets, x, 2.35, w, self.H - 3.1, 20, self.c["text"], self.c["accent"])
        self._chrome(slide, i, total, bg)

    def _l_quote(self, slide, i, total, s, deck, n):
        bg = self.c["surface"]
        self._bg(slide, bg, force=True)
        fg = text_on(bg)
        self._text(slide, i, M, 0.2, 2.5, 2.5, "“", 160, self.c["accent"], heading=True)
        self._text(slide, i, M + 0.6, 1.8, self.W - 2 * M - 1.2, 3.2, s.quote or s.title, 30, fg, italic=True,
                   heading=True, anchor="middle")
        if s.attribution:
            self._text(slide, i, M + 0.6, 5.2, self.W - 2 * M - 1.2, 0.6, "— " + s.attribution, 18, mix(fg, bg, 0.4))
        self._chrome(slide, i, total, bg)
